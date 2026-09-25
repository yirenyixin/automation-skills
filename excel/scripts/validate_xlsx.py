"""报告 XLSX 中未替换的占位符、已保存的错误值和空工作表。"""
from __future__ import annotations
import argparse, json, re
from pathlib import Path
from openpyxl import load_workbook
ERRORS = {'#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#N/A', '#NUM!', '#NULL!', '#SPILL!', '#CALC!'}
def main():
    p = argparse.ArgumentParser(description=__doc__); p.add_argument('source', type=Path); p.add_argument('--output', required=True, type=Path); a = p.parse_args(); wb = load_workbook(a.source, data_only=False)
    placeholders, errors, empty = [], [], []
    for ws in wb.worksheets:
        if ws.max_row == 1 and ws.max_column == 1 and ws['A1'].value is None: empty.append(ws.title)
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str): placeholders.extend({'sheet': ws.title, 'cell': cell.coordinate, 'marker': x} for x in re.findall(r'\{\{[^{}]+\}\}', cell.value))
                if cell.value in ERRORS: errors.append({'sheet': ws.title, 'cell': cell.coordinate, 'error': cell.value})
    report = {'source': str(a.source), 'sheet_count': len(wb.worksheets), 'unresolved_placeholders': placeholders, 'saved_error_values': errors, 'empty_sheets': empty, 'valid': not placeholders and not errors}
    a.output.parent.mkdir(parents=True, exist_ok=True); a.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8'); print(json.dumps(report, ensure_ascii=False))
if __name__ == '__main__': main()
