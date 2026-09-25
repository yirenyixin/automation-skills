"""替换复制后 XLSX 模板中的标量 {{字段}} 占位符。"""
from __future__ import annotations
import argparse, json
from pathlib import Path
from openpyxl import load_workbook

def main():
    p = argparse.ArgumentParser(description=__doc__); p.add_argument('template', type=Path); p.add_argument('data', type=Path); p.add_argument('--output', required=True, type=Path); a = p.parse_args()
    values = json.loads(a.data.read_text(encoding='utf-8'))
    if not isinstance(values, dict) or any(isinstance(v, (dict, list)) for v in values.values()): p.error('数据必须是值为标量的 JSON 对象')
    wb = load_workbook(a.template)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    text = cell.value
                    for key, value in values.items(): text = text.replace('{{' + key + '}}', '' if value is None else str(value))
                    cell.value = text
    a.output.parent.mkdir(parents=True, exist_ok=True); wb.save(a.output)
if __name__ == '__main__': main()
