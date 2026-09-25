"""从受限 JSON IR 创建 XLSX 工作簿。"""
from __future__ import annotations
import argparse, json
from datetime import date, datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

def value(item):
    if isinstance(item, dict) and item.get('type') == 'date': return datetime.fromisoformat(item['value']).date()
    return item

def main():
    p = argparse.ArgumentParser(description=__doc__); p.add_argument('input', type=Path); p.add_argument('--output', required=True, type=Path); a = p.parse_args(); data = json.loads(a.input.read_text(encoding='utf-8'))
    sheets = data.get('sheets', [])
    if not sheets: p.error('至少需要一个工作表')
    wb = Workbook(); wb.remove(wb.active); names = set()
    for spec in sheets:
        name = str(spec.get('name', ''))
        if not name or len(name) > 31 or name in names: p.error(f'无效或重复的工作表名：{name!r}')
        names.add(name); ws = wb.create_sheet(name); headers = spec.get('headers', []); rows = spec.get('rows', [])
        if headers:
            ws.append([value(x) for x in headers])
            for cell in ws[1]: cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F4E78')
        for row in rows: ws.append([value(x) for x in row])
        if spec.get('freeze_panes'): ws.freeze_panes = spec['freeze_panes']
        for column in ws.columns:
            width = min(50, max(10, max(len(str(cell.value or '')) for cell in column) + 2)); ws.column_dimensions[column[0].column_letter].width = width
    a.output.parent.mkdir(parents=True, exist_ok=True); wb.save(a.output)
if __name__ == '__main__': main()
