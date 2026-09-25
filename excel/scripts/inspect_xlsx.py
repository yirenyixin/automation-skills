"""将 XLSX 的工作表、单元格值和公式提取为 JSON。"""
from __future__ import annotations
import argparse, json
from pathlib import Path
from openpyxl import load_workbook

def main():
    p = argparse.ArgumentParser(description=__doc__); p.add_argument('source', type=Path); p.add_argument('--output', required=True, type=Path); a = p.parse_args()
    wb = load_workbook(a.source, data_only=False, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        cells = [{'coordinate': cell.coordinate, 'value': cell.value, 'data_type': cell.data_type, 'number_format': cell.number_format} for row in ws.iter_rows() for cell in row if cell.value is not None]
        sheets.append({'name': ws.title, 'max_row': ws.max_row, 'max_column': ws.max_column, 'freeze_panes': str(ws.freeze_panes or ''), 'cells': cells})
    a.output.parent.mkdir(parents=True, exist_ok=True); a.output.write_text(json.dumps({'sheets': sheets}, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
if __name__ == '__main__': main()
