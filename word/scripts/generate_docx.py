"""Generate DOCX headings, paragraphs, and tables from constrained JSON IR."""
from __future__ import annotations
import argparse, json
from pathlib import Path
from docx import Document

def main():
    p = argparse.ArgumentParser(description=__doc__); p.add_argument('input', type=Path); p.add_argument('--output', required=True, type=Path); a = p.parse_args()
    data = json.loads(a.input.read_text(encoding='utf-8')); d = Document()
    if data.get('title'): d.add_heading(str(data['title']), level=0)
    for block in data.get('blocks', []):
        kind = block.get('type')
        if kind == 'heading': d.add_heading(str(block.get('text', '')), level=max(1, min(9, int(block.get('level', 1)))))
        elif kind == 'paragraph': d.add_paragraph(str(block.get('text', '')))
        elif kind == 'table':
            headers, rows = block.get('headers', []), block.get('rows', [])
            table = d.add_table(rows=1 if headers else 0, cols=len(headers), style='Table Grid')
            if headers:
                for cell, value in zip(table.rows[0].cells, headers): cell.text = str(value)
            for row in rows:
                for cell, value in zip(table.add_row().cells, row): cell.text = str(value)
        else: p.error(f'unsupported block type: {kind!r}')
    a.output.parent.mkdir(parents=True, exist_ok=True); d.save(a.output)
if __name__ == '__main__': main()
