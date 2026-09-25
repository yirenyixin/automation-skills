"""填充单个 run 内的 {{field}}，保留段落和 run 格式。"""
from __future__ import annotations
import argparse, json, re
from pathlib import Path
from docx import Document

def all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells: yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs

def replace_in_paragraph(paragraph, values):
    combined = ''.join(run.text for run in paragraph.runs)
    for marker in re.findall(r'\{\{[^{}]+\}\}', combined):
        if not any(marker in run.text for run in paragraph.runs):
            raise ValueError(f'占位符 {marker} 跨越多个 run；请将其置于单一 run 以保留格式')
    for run in paragraph.runs:
        for key, value in values.items(): run.text = run.text.replace('{{' + key + '}}', '' if value is None else str(value))

def main():
    p = argparse.ArgumentParser(description=__doc__); p.add_argument('template', type=Path); p.add_argument('data', type=Path); p.add_argument('--output', required=True, type=Path); a = p.parse_args()
    if a.template.resolve() == a.output.resolve(): p.error('输出路径不得与源模板相同')
    values = json.loads(a.data.read_text(encoding='utf-8'))
    if not isinstance(values, dict) or any(isinstance(v, (dict, list)) for v in values.values()): p.error('data must be a JSON object with scalar values')
    d = Document(a.template)
    try:
        for paragraph in all_paragraphs(d): replace_in_paragraph(paragraph, values)
    except ValueError as exc: p.error(str(exc))
    a.output.parent.mkdir(parents=True, exist_ok=True); d.save(a.output)
if __name__ == '__main__': main()
