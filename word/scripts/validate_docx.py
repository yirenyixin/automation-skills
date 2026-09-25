"""Check a DOCX for unresolved placeholders and basic structural signals."""
from __future__ import annotations
import argparse, json, re
from pathlib import Path
from docx import Document

def main():
    p = argparse.ArgumentParser(description=__doc__); p.add_argument('source', type=Path); p.add_argument('--output', required=True, type=Path); a = p.parse_args()
    d = Document(a.source); texts = [x.text for x in d.paragraphs] + [cell.text for table in d.tables for row in table.rows for cell in row.cells]
    unresolved = sorted({m.group(0) for text in texts for m in re.finditer(r'\{\{[^{}]+\}\}', text)})
    report = {'source': str(a.source), 'paragraph_count': len(d.paragraphs), 'table_count': len(d.tables), 'unresolved_placeholders': unresolved, 'empty_text_nodes': sum(not x.strip() for x in texts), 'valid': not unresolved}
    a.output.parent.mkdir(parents=True, exist_ok=True); a.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8'); print(json.dumps(report, ensure_ascii=False))
if __name__ == '__main__': main()
