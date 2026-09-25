"""Extract ordinary paragraphs, tables, and core metadata from DOCX to JSON."""
from __future__ import annotations
import argparse, json
from pathlib import Path
from docx import Document

def main():
    p = argparse.ArgumentParser(description=__doc__); p.add_argument('source', type=Path); p.add_argument('--output', required=True, type=Path); a = p.parse_args()
    d = Document(a.source)
    result = {'metadata': {k: str(getattr(d.core_properties, k) or '') for k in ('title', 'author', 'subject', 'keywords')}, 'paragraphs': [{'index': i, 'text': x.text, 'style': x.style.name} for i, x in enumerate(d.paragraphs)], 'tables': [[[cell.text for cell in row.cells] for row in table.rows] for table in d.tables]}
    a.output.parent.mkdir(parents=True, exist_ok=True); a.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
if __name__ == '__main__': main()
